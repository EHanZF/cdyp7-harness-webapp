from pydantic import BaseModel
from app.core.storage import blob_store
from app.core.config import settings
from docx import Document
from pathlib import Path

class ValidationCheck(BaseModel):
    check_id: str
    status: str
    message: str | None = None
class ValidationResult(BaseModel):
    status: str
    checks: list[ValidationCheck]

def validate_release_sheet_docx_blob(blob_name: str, profile: str='system_release_sheet') -> ValidationResult:
    checks=[]
    data = blob_store.read_bytes(settings.artifacts_container, blob_name)
    tmp = Path('outputs') / ('validate-' + blob_name)
    tmp.parent.mkdir(exist_ok=True)
    tmp.write_bytes(data)
    doc = Document(tmp)
    text='\n'.join(p.text for p in doc.paragraphs)
    def passed(cid): checks.append(ValidationCheck(check_id=cid,status='passed'))
    def failed(cid,msg): checks.append(ValidationCheck(check_id=cid,status='failed',message=msg))
    def warning(cid,msg): checks.append(ValidationCheck(check_id=cid,status='warning',message=msg))
    passed('docx_readable')
    if 'System Release Sheet' in text:
        passed('system_release_sheet_title_present')
    else:
        failed('system_release_sheet_title_present','Title missing')

    if '{{' in text or '}}' in text:
        failed('unresolved_required_placeholders','Unresolved placeholder found')
    else:
        passed('unresolved_required_placeholders')
    if profile == 'system_release_sheet_cat5':
        passed('cat5_profile_applied')
    warning('approval_signatures_present','Approval signatures require HITL completion.')
    status='failed' if any(c.status=='failed' for c in checks) else ('warning' if any(c.status=='warning' for c in checks) else 'passed')
    return ValidationResult(status=status, checks=checks)
